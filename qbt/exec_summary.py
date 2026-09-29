"""Executive summary: a one-page CIO report at the front of the workbook,
and professional formatting for the sheets behind it.

The summary is a printable page -- portrait, fitted to one sheet of paper
-- that condenses the run into what a CIO reads first: headline figures
against the benchmark, trailing and calendar returns, risk statistics, a
growth chart, behaviour in stress episodes and market regimes, robustness,
key observations, the analyst's own comments, and every assumption the
figures depend on (costs, fees, execution, data).

The observations are generated from the numbers by fixed rules -- they
state what the figures show and never add an opinion the figures do not
support. The analyst's comments are where judgment goes.

Everything is written with openpyxl into the workbook the export already
builds, so it adds a fraction of a second, and only when the file is
actually downloaded.
"""
from __future__ import annotations

from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .brand import PRINT, RED

SHEET = "Executive Summary"
_DATA = "_Chart Data"

INK = PRINT["text"].lstrip("#")
MUTED = PRINT["muted"].lstrip("#")
RULE = PRINT["rule"].lstrip("#")
SAND = PRINT["sand"].lstrip("#")
GAIN = PRINT["gain"].lstrip("#")
LOSS = PRINT["loss"].lstrip("#")
BRAND_RED = RED.lstrip("#")
FONT = "Calibri"

PCT = '0.0%;-0.0%;0.0%'
PCT2 = '0.00%;-0.00%;0.00%'
NUM = '0.00'
MONEY = '$#,##0'

# Statistics shown as percentages.
_PCT_KEYS = {"Total Return", "CAGR", "Volatility", "Max Drawdown", "VaR 95% (daily)",
             "CVaR 95% (daily)", "% Positive Months", "Best Month", "Worst Month",
             "Alpha (ann.)", "Tracking Error", "Average Exposure"}

# Two blocks side by side: B..E on the left, G..J on the right.
LEFT, RIGHT = 2, 7
WIDTHS = {1: 1.5, 2: 20, 3: 11, 4: 11, 5: 11, 6: 2.5, 7: 20, 8: 11, 9: 11, 10: 11, 11: 1.5}


# ----------------------------------------------------------------------
def _f(size=9, bold=False, color=INK, italic=False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def _isnum(v) -> bool:
    return isinstance(v, (int, float, np.integer, np.floating)) and np.isfinite(v)


class _Page:
    def __init__(self, ws):
        self.ws = ws
        self.row = 1

    def cell(self, r, c, v=None, fmt=None, font=None, align=None, fill=None,
             signed=False):
        cell = self.ws.cell(row=r, column=c)
        if v is not None and not (isinstance(v, float) and not np.isfinite(v)):
            cell.value = v.item() if isinstance(v, np.generic) else v
        elif v is not None:
            cell.value = "n/a"
        if fmt and _isnum(v):
            cell.number_format = fmt
        f = font or _f()
        if signed and _isnum(v) and v != 0:
            f = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic,
                     color=GAIN if v > 0 else LOSS)
        cell.font = f
        cell.alignment = align or Alignment(
            horizontal="right" if _isnum(v) or v == "n/a" else "left",
            vertical="center")
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        return cell

    def rule(self, r, c0, c1, color=RULE, style="thin"):
        for c in range(c0, c1 + 1):
            cell = self.ws.cell(row=r, column=c)
            cell.border = Border(bottom=Side(style=style, color=color))

    def heading(self, r, c0, c1, text):
        self.cell(r, c0, text.upper(), font=_f(8.5, True, BRAND_RED))
        self.rule(r, c0, c1, BRAND_RED, "thin")

    def table(self, r, c0, header: Sequence[str], rows: List[Sequence[Any]],
              fmts: Sequence[Optional[str]], signed: Sequence[bool] = ()) -> int:
        """A compact table; returns the next free row."""
        for j, h in enumerate(header):
            self.cell(r, c0 + j, h, font=_f(8, True, MUTED),
                      align=Alignment(horizontal="left" if j == 0 else "right",
                                      vertical="center"))
        self.rule(r, c0, c0 + len(header) - 1)
        r += 1
        for k, row in enumerate(rows):
            for j, v in enumerate(row):
                self.cell(r, c0 + j, v, fmt=fmts[j] if j < len(fmts) else None,
                          signed=bool(signed[j]) if j < len(signed) else False,
                          fill=SAND if k % 2 == 1 else None)
            r += 1
        return r

    def para(self, r, c0, c1, text, font=None, height=None):
        self.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        self.cell(r, c0, text, font=font or _f(9),
                  align=Alignment(wrap_text=True, vertical="top", horizontal="left"))
        if height:
            self.ws.row_dimensions[r].height = height


def _lines_needed(text: str, chars_per_line: int = 125) -> int:
    return max(1, int(np.ceil(len(text) / chars_per_line)))


# ----------------------------------------------------------------------
# Observations: stated from the numbers, by fixed rules.
def _pct(v, d=1) -> str:
    return "n/a" if not _isnum(v) else f"{v * 100:+.{d}f}%".replace("+", "") if v >= 0 \
        else f"{v * 100:.{d}f}%"


def _pts(v) -> str:
    return "n/a" if not _isnum(v) else f"{v * 100:+.1f} pts"


def observations(d: Dict[str, Any]) -> List[str]:
    st, bs = d.get("stats", {}) or {}, d.get("bench_stats", {}) or {}
    bl = d.get("bench_label")
    out: List[str] = []
    yrs = d.get("years")
    c, bc = st.get("CAGR"), bs.get("CAGR")
    if bl and _isnum(c) and _isnum(bc):
        out.append(f"Compounded {_pct(c)} a year over {yrs:.1f} years, against "
                   f"{_pct(bc)} for {bl}: {_pts(c - bc)} a year "
                   f"{'ahead of' if c >= bc else 'behind'} the benchmark.")
    elif _isnum(c):
        out.append(f"Compounded {_pct(c)} a year over {yrs:.1f} years.")
    v, bv, dd, bdd = (st.get("Volatility"), bs.get("Volatility"),
                      st.get("Max Drawdown"), bs.get("Max Drawdown"))
    if bl and all(_isnum(x) for x in (v, bv, dd, bdd)):
        out.append(f"Volatility of {_pct(v)} against {_pct(bv)}, and a deepest "
                   f"drawdown of {_pct(dd)} against {_pct(bdd)}"
                   + (": less risk taken for the return." if v < bv and dd > bdd else
                      ": more risk taken than the benchmark." if v > bv and dd < bdd else
                      "."))
    s, bsr, ir = st.get("Sharpe"), bs.get("Sharpe"), st.get("Information Ratio")
    if _isnum(s):
        txt = f"Sharpe ratio {s:.2f}" + (f" against {bsr:.2f} for the benchmark"
                                         if bl and _isnum(bsr) else "")
        if _isnum(ir):
            txt += f"; information ratio {ir:.2f}"
        out.append(txt + ".")
    dds = d.get("drawdowns")
    if dds is not None and not dds.empty:
        top = dds.iloc[0]
        rec = str(top.get("Recovery"))
        out.append(f"Worst episode: {_pct(top['Drawdown'])} from {top['Start']} to "
                   f"{top['Trough']}, " + ("not yet recovered." if rec == "ongoing"
                                           else f"recovered by {rec}."))

    ev = d.get("stress")
    if ev is not None and not ev.empty:
        cov = ev[ev["Coverage"].isin(["Full", "Partial"])].dropna(subset=["Return"])
        if len(cov):
            pos = int((cov["Return"] > 0).sum())
            txt = f"Stress episodes: positive in {pos} of {len(cov)}"
            ex = cov["Excess vs Benchmark"].dropna()
            if len(ex):
                txt += f", ahead of the benchmark in {int((ex > 0).sum())} of {len(ex)}"
            w = cov.loc[cov["Return"].idxmin()]
            txt += f"; worst was {w['Period']} at {_pct(w['Return'])}"
            if _isnum(w.get("Benchmark Return")):
                txt += f" ({_pct(w['Benchmark Return'])} for the benchmark)"
            out.append(txt + ".")

    rg = d.get("regimes")
    if rg is not None and not rg.empty:
        mine = rg[rg["Series"] == d.get("label")].dropna(subset=["Ann. return"])
        mine = mine[mine["% of time"] >= 0.05]
        if len(mine):
            best, worst = mine.loc[mine["Ann. return"].idxmax()], mine.loc[mine["Ann. return"].idxmin()]
            txt = (f"Across market regimes, strongest in {best['Regime']} "
                   f"({_pct(best['Ann. return'])} a year) and weakest in "
                   f"{worst['Regime']} ({_pct(worst['Ann. return'])})")
            if "Excess ann. return" in mine and mine["Excess ann. return"].notna().any():
                e = mine.loc[mine["Excess ann. return"].idxmax()]
                txt += (f"; the largest edge over the benchmark came in "
                        f"{e['Regime']} ({_pts(e['Excess ann. return'])})")
            out.append(txt + ".")

    wf = d.get("walk_forward")
    if wf is not None and not wf.empty and "Sharpe" in wf:
        sh = wf["Sharpe"].dropna()
        if len(sh) > 1:
            disp = float(sh.std(ddof=1))
            out.append(f"Sub-period Sharpe ranged from {sh.min():.2f} to {sh.max():.2f} "
                       f"across {len(sh)} folds"
                       + ("; the result depends noticeably on the market regime."
                          if disp > 0.6 else "; reasonably stable over time."))
    ios = d.get("in_out")
    if ios is not None and len(ios) == 2 and "Sharpe" in ios:
        i_, o_ = ios["Sharpe"].iloc[0], ios["Sharpe"].iloc[1]
        if _isnum(i_) and _isnum(o_):
            out.append(f"Out-of-sample Sharpe {o_:.2f} against {i_:.2f} in-sample"
                       + (": no sign of decay." if o_ >= 0.8 * i_ else
                          ": weaker out of sample, a caution against overfitting."))
    mc = d.get("monte_carlo")
    if isinstance(mc, dict) and _isnum(mc.get("median_cagr")):
        out.append(f"Resampled {mc.get('n', 'many')} ways, the median CAGR is "
                   f"{_pct(mc['median_cagr'])}, with a {mc['prob_loss'] * 100:.0f}% "
                   f"chance of ending below the start and "
                   f"{mc['prob_dd_20'] * 100:.0f}% of a drawdown beyond 20%.")
    cpa, fee, to = d.get("cost_pa"), d.get("fee_pa"), st.get("Annual Turnover")
    if _isnum(cpa):
        txt = f"Trading costs took {_pct(cpa, 2)} a year"
        if _isnum(to):
            txt += f" on turnover of {to:.1f}x"
        if _isnum(fee) and fee > 0:
            txt += f"; the management fee a further {_pct(fee, 2)}"
        out.append(txt + ".")
    tax = d.get("tax")
    if tax:
        out.append(f"Tax Cost Ratio of {_pct(tax['tcr'], 2)} a year "
                   f"({tax['label'].lower()}) for a Canadian taxable account"
                   + (f", against {_pct(tax['bench_tcr'], 2)} for the benchmark"
                      if _isnum(tax.get('bench_tcr')) else "") + ".")
    return out


# ----------------------------------------------------------------------
def write_summary(wb, d: Dict[str, Any]) -> None:
    """Adds the one-page summary as the first sheet of `wb`."""
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = 12
    ws.sheet_format.customHeight = True
    for c, w in WIDTHS.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    p = _Page(ws)
    label, bl = d.get("label", "Strategy"), d.get("bench_label")
    st, bs = d.get("stats", {}) or {}, d.get("bench_stats", {}) or {}

    # --- Masthead --------------------------------------------------------
    ws.row_dimensions[1].height = 6
    ws.merge_cells("B2:H2")
    p.cell(2, 2, d.get("title", label), font=_f(16, True))
    ws.merge_cells("I2:J2")
    p.cell(2, 9, "EXECUTIVE SUMMARY", font=_f(8.5, True, BRAND_RED),
           align=Alignment(horizontal="right", vertical="center"))
    ws.row_dimensions[2].height = 24
    ws.merge_cells("B3:J3")
    p.cell(3, 2, d.get("subtitle", ""), font=_f(8.5, color=MUTED))
    p.rule(3, 2, 10, BRAND_RED, "medium")

    # --- KPI tiles -------------------------------------------------------
    r = 5
    tiles = [("CAGR", PCT), ("Volatility", PCT), ("Sharpe", NUM), ("Sortino", NUM),
             ("Max Drawdown", PCT), ("Calmar", NUM), ("Beta", NUM),
             ("Information Ratio", NUM)]
    cols = [2, 3, 4, 5, 7, 8, 9, 10]
    for (k, fmt), c in zip(tiles, cols):
        p.cell(r, c, k.upper() if k != "Information Ratio" else "INFO RATIO",
               font=_f(7, True, MUTED), align=Alignment(horizontal="center"),
               fill=SAND)
        v = st.get(k)
        p.cell(r + 1, c, v, fmt=fmt, font=_f(14, True,
               (GAIN if _isnum(v) and v > 0 else LOSS if _isnum(v) and v < 0 else INK)
               if k in ("CAGR", "Sharpe", "Sortino", "Calmar", "Information Ratio") else INK),
               align=Alignment(horizontal="center", vertical="center"), fill=SAND)
        bv = bs.get(k) if bl else None
        txt = "" if not bl or k in ("Beta", "Information Ratio") else (
            "bench " + ("n/a" if not _isnum(bv) else
                        f"{bv * 100:.1f}%" if fmt == PCT else f"{bv:.2f}"))
        p.cell(r + 2, c, txt, font=_f(7, color=MUTED),
               align=Alignment(horizontal="center"), fill=SAND)
    ws.row_dimensions[r + 1].height = 22
    r += 4

    # --- Trailing | Calendar years -----------------------------------------
    hdr3 = ["", "Strategy", "Benchmark" if bl else "", "Excess" if bl else ""]
    p.heading(r, LEFT, LEFT + 3, "Trailing returns")
    p.heading(r, RIGHT, RIGHT + 3, "Calendar years")
    r += 1
    left_end = right_end = r
    tr = d.get("trailing")
    if tr is not None and not tr.empty:
        want = ["YTD", "1Y", "3Y", "5Y", "10Y", "Since inception"]
        t = tr[tr["Period"].isin(want)]
        vcols = [c for c in t.columns if c not in ("Period", "Annualized", "From", "To", "Excess")]
        rows = []
        for _, x in t.iterrows():
            ann = " (ann.)" if str(x.get("Annualized", "")).lower().startswith("y") else ""
            rows.append([f"{x['Period']}{ann}", x.get(vcols[0]),
                         x.get(vcols[1]) if len(vcols) > 1 else None,
                         x.get("Excess") if "Excess" in t else None])
        left_end = p.table(r, LEFT, hdr3, rows, [None, PCT, PCT, PCT],
                           [False, True, True, True])
    cal = d.get("calendar")
    if cal is not None and not cal.empty:
        vcols = [c for c in cal.columns if c not in ("Year", "Partial", "Excess")]
        rows = [[f"{int(x['Year'])}{' (partial)' if str(x.get('Partial', '')).strip() else ''}",
                 x[vcols[0]], x[vcols[1]] if len(vcols) > 1 else None,
                 x.get("Excess") if "Excess" in cal else None]
                for _, x in cal.tail(6).iterrows()]
        right_end = p.table(r, RIGHT, hdr3, rows, [None, PCT, PCT, PCT],
                            [False, True, True, True])
    r = max(left_end, right_end) + 1

    # --- Growth chart ----------------------------------------------------
    chart_rows: set = set()
    curve = d.get("curve")
    if curve is not None and not curve.empty:
        p.heading(r, LEFT, 10, f"Growth of {d.get('capital_label', '$100,000')}")
        r += 1
        _chart(wb, ws, curve, r)
        chart_rows = set(range(r, r + 11))
        r += 11

    # --- Risk | Stress -----------------------------------------------------
    p.heading(r, LEFT, LEFT + 3, "Risk and quantitative profile")
    p.heading(r, RIGHT, RIGHT + 3, "Stress episodes (worst five)")
    r += 1
    risk = [("Beta", NUM), ("Alpha (ann.)", PCT), ("Tracking Error", PCT),
            ("Information Ratio", NUM), ("VaR 95% (daily)", PCT2),
            ("CVaR 95% (daily)", PCT2), ("Annual Turnover", NUM)]
    rows = [(k, st.get(k), bs.get(k) if bl else None, fmt) for k, fmt in risk if k in st]
    for j, h in enumerate(["", "Strategy", "Benchmark" if bl else ""]):
        p.cell(r, LEFT + j, h, font=_f(8, True, MUTED),
               align=Alignment(horizontal="left" if j == 0 else "right"))
    p.rule(r, LEFT, LEFT + 3)
    rr = r + 1
    for k_i, (k, a, b, fmt) in enumerate(rows):
        fill = SAND if k_i % 2 else None
        p.cell(rr, LEFT, k, fill=fill)
        p.cell(rr, LEFT + 1, a, fmt=fmt, fill=fill)
        p.cell(rr, LEFT + 2, b if b is not None else "", fmt=fmt, fill=fill)
        p.cell(rr, LEFT + 3, None, fill=fill)
        rr += 1
    left_end, right_end = rr, r
    ev = d.get("stress")
    if ev is not None and not ev.empty:
        cov = ev[ev["Coverage"].isin(["Full", "Partial"])].dropna(subset=["Return"])
        rows = [[str(x["Period"])[:24], x["Return"], x.get("Benchmark Return"),
                 x.get("Excess vs Benchmark")]
                for _, x in cov.sort_values("Return").head(5).iterrows()]
        if rows:
            right_end = p.table(r, RIGHT, ["", "Strategy", "Benchmark" if bl else "",
                                           "Excess" if bl else ""],
                                rows, [None, PCT, PCT, PCT], [False, True, True, True])
    if right_end == r:
        p.cell(r, RIGHT, "No stress episode falls inside the period.",
               font=_f(8.5, italic=True, color=MUTED))
        right_end = r + 1
    r = max(left_end, right_end) + 1

    # --- Regimes | Robustness ----------------------------------------------
    p.heading(r, LEFT, LEFT + 3, "Market regimes (annualized)")
    p.heading(r, RIGHT, RIGHT + 3, "Robustness")
    r += 1
    left_end = right_end = r
    rg = d.get("regimes")
    if rg is not None and not rg.empty:
        piv = rg.pivot_table(index=["Dimension", "Regime"], columns="Series",
                             values="Ann. return", sort=False)
        rows = []
        for dim in ("Volatility (VIX)", "Economy (NBER recessions)"):
            if dim not in piv.index.get_level_values(0):
                continue
            for reg, x in piv.loc[dim].iterrows():
                a, b = x.get(label), x.get(bl) if bl else None
                rows.append([str(reg)[:24], a, b,
                             a - b if _isnum(a) and _isnum(b) else None])
        if rows:
            left_end = p.table(r, LEFT, hdr3, rows, [None, PCT, PCT, PCT],
                               [False, True, True, True])
    if left_end == r:
        p.cell(r, LEFT, "Regime data not included in this export.",
               font=_f(8.5, italic=True, color=MUTED))
        left_end = r + 1
    rob: List[Tuple[str, Any, str]] = []
    wf = d.get("walk_forward")
    if wf is not None and not wf.empty and "Sharpe" in wf:
        rob += [("Sub-period Sharpe, lowest", wf["Sharpe"].min(), NUM),
                ("Sub-period Sharpe, highest", wf["Sharpe"].max(), NUM)]
    ios = d.get("in_out")
    if ios is not None and len(ios) == 2:
        rob += [("In-sample Sharpe", ios["Sharpe"].iloc[0], NUM),
                ("Out-of-sample Sharpe", ios["Sharpe"].iloc[1], NUM)]
    cs = d.get("cost")
    if cs is not None and not cs.empty:
        hit = cs[cs[cs.columns[0]] == 50]
        if len(hit):
            rob.append(("CAGR at 50 bps per trade", hit["CAGR"].iloc[0], PCT))
    mc = d.get("monte_carlo")
    if isinstance(mc, dict) and _isnum(mc.get("median_cagr")):
        rob += [("Monte Carlo median CAGR", mc["median_cagr"], PCT),
                ("Monte Carlo, probability of loss", mc["prob_loss"], PCT)]
    rr = r
    for k_i, (k, v, fmt) in enumerate(rob[:6]):
        fill = SAND if k_i % 2 else None
        ws.merge_cells(start_row=rr, start_column=RIGHT, end_row=rr, end_column=RIGHT + 2)
        p.cell(rr, RIGHT, k, fill=fill)
        p.cell(rr, RIGHT + 3, v, fmt=fmt, fill=fill)
        rr += 1
    if not rob:
        p.cell(rr, RIGHT, "Not included in this export.",
               font=_f(8.5, italic=True, color=MUTED))
        rr += 1
    right_end = rr
    r = max(left_end, right_end) + 1

    # --- Observations -----------------------------------------------------
    _order = ("Compounded", "Volatility", "Sharpe", "Stress", "Across",
              "Sub-period", "Out-of-sample", "Trading costs", "Tax Cost",
              "Worst episode", "Resampled")
    obs = sorted(observations(d), key=lambda o: next(
        (i for i, k in enumerate(_order) if o.startswith(k)), len(_order)))[:9]
    if obs:
        p.heading(r, LEFT, 10, "Key observations")
        r += 1
        for o in obs:
            txt = f"•  {o}"
            p.para(r, LEFT, 10, txt, height=12 * _lines_needed(txt, 135))
            r += 1
        r += 1

    comments = (d.get("comments") or "").strip()
    p.heading(r, LEFT, 10, "Comments for the CIO")
    r += 1
    if comments:
        for para in [x.strip() for x in comments.split("\n") if x.strip()]:
            p.para(r, LEFT, 10, para, height=12.5 * _lines_needed(para))
            r += 1
    else:
        p.para(r, LEFT, 10, "", height=40)
        ws.cell(row=r, column=LEFT).border = Border()
        r += 1
    r += 1

    # --- Assumptions -------------------------------------------------------
    p.heading(r, LEFT, 10, "Assumptions, costs and fees")
    r += 1
    items = list((d.get("assumptions") or {}).items())
    half = (len(items) + 1) // 2
    for k_i, (k, v) in enumerate(items[:half]):
        fill = SAND if k_i % 2 else None
        p.cell(r + k_i, LEFT, k, font=_f(8.5, color=MUTED), fill=fill)
        ws.merge_cells(start_row=r + k_i, start_column=LEFT + 1, end_row=r + k_i,
                       end_column=LEFT + 3)
        p.cell(r + k_i, LEFT + 1, str(v), font=_f(8.5), fill=fill,
               align=Alignment(horizontal="left", vertical="center"))
    for k_i, (k, v) in enumerate(items[half:]):
        fill = SAND if k_i % 2 else None
        p.cell(r + k_i, RIGHT, k, font=_f(8.5, color=MUTED), fill=fill)
        ws.merge_cells(start_row=r + k_i, start_column=RIGHT + 1, end_row=r + k_i,
                       end_column=RIGHT + 3)
        p.cell(r + k_i, RIGHT + 1, str(v), font=_f(8.5), fill=fill,
               align=Alignment(horizontal="left", vertical="center"))
    r += half + 1

    note = ("Backtested, hypothetical results computed from historical data; "
            "they do not represent actual trading and are not a guarantee of "
            "future performance. Returns are net of the trading costs and fees "
            "listed above and before taxes unless stated. Figures are drawn from "
            "the sheets of this workbook, which carry every underlying table.")
    p.para(r, LEFT, 10, note, font=_f(7.5, italic=True, color=MUTED),
           height=11 * _lines_needed(note, 140))
    last = r

    # Blank rows between sections become thin gaps, which is most of what
    # keeps the page on one sheet at a readable scale.
    for rr in range(4, last):
        if rr in chart_rows or ws.row_dimensions[rr].height:
            continue
        if all(ws.cell(row=rr, column=c).value in (None, "") for c in range(1, 12)):
            ws.row_dimensions[rr].height = 5

    # --- Print: one portrait page -----------------------------------------
    ws.print_area = f"A1:K{last}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.35, right=0.35, top=0.4, bottom=0.4)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = "&8" + d.get("footer", "")
    ws.oddFooter.right.text = "&8Page &P of &N"
    wb.active = 0


def _chart(wb, ws, curve: pd.DataFrame, anchor_row: int) -> None:
    """Native Excel line chart of month-end values, from a hidden sheet."""
    if _DATA in wb.sheetnames:
        del wb[_DATA]
    ds = wb.create_sheet(_DATA)
    ds.sheet_state = "hidden"
    ds.cell(row=1, column=1, value="Date")
    for j, c in enumerate(curve.columns, start=2):
        ds.cell(row=1, column=j, value=str(c))
    for i, (dt, row) in enumerate(curve.iterrows(), start=2):
        ds.cell(row=i, column=1, value=dt.to_pydatetime()).number_format = "yyyy-mm"
        for j, v in enumerate(row.values, start=2):
            if _isnum(v):
                ds.cell(row=i, column=j, value=float(v))
    n = len(curve) + 1
    ch = LineChart()
    ch.height, ch.width = 5.6, 25.5
    ch.legend.position = "t"
    ch.y_axis.number_format = MONEY
    ch.y_axis.majorGridlines.spPr = None
    ch.x_axis.number_format = "yyyy"
    ch.x_axis.majorTimeUnit = "years"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.add_data(Reference(ds, min_col=2, max_col=1 + len(curve.columns),
                          min_row=1, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(ds, min_col=1, min_row=2, max_row=n))
    colors = [BRAND_RED, "3F3A34", "1F6F5C"]
    for s, col in zip(ch.series, colors):
        s.graphicalProperties.line.solidFill = col
        s.graphicalProperties.line.width = 19050
        s.smooth = False
    ws.add_chart(ch, f"B{anchor_row}")
